"""Build reproducible, expert-reference-relative features for pilot model research."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from form_matching import compare_form_images
from multimodal import extract_stroke_features
from stroke_matching import compare_stroke_sequences


FEATURE_SCHEMA_VERSION = 'pilot-v1'
REQUIRED_MANIFEST_COLUMNS = {
    'sample_id', 'relative_path', 'writer_code', 'target', 'condition', 'binary_label',
    'quality_status', 'schema_status', 'training_eligibility',
}


class TrainingDataError(ValueError):
    """Raised when a provisional training record is unsafe or incomplete."""


def _inside(root: Path, candidate: Path) -> Path:
    resolved_root, resolved_candidate = root.resolve(), candidate.resolve()
    if resolved_root != resolved_candidate and resolved_root not in resolved_candidate.parents:
        raise TrainingDataError('Manifest path escapes the project directory')
    return resolved_candidate


def load_cleaned_manifest(path: Path) -> list[dict]:
    """Load eligible rows from the cleaned workbook without changing the raw recordings."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        try:
            sheet = workbook['Sheet1']
        except KeyError as error:
            raise TrainingDataError('Cleaned manifest must contain Sheet1') from error
        header_index, header_row = next(
            (
                (index, row) for index, row in enumerate(sheet.iter_rows(max_row=10, values_only=True), start=1)
                if REQUIRED_MANIFEST_COLUMNS.issubset(set(row))
            ),
            (None, None),
        )
        if header_row is None:
            raise TrainingDataError('Cleaned manifest is missing required columns')
        headers = [str(value) if value is not None else '' for value in header_row]
        rows = []
        for values in sheet.iter_rows(min_row=header_index + 1, values_only=True):
            record = {header: value for header, value in zip(headers, values)}
            if not record['sample_id']:
                continue
            if record['training_eligibility'] != 'eligible':
                continue
            if record['quality_status'] != 'accepted' or record['schema_status'] != 'valid':
                raise TrainingDataError(f"Eligible sample {record['sample_id']} is not accepted and valid")
            rows.append({column: str(record[column]) for column in REQUIRED_MANIFEST_COLUMNS})
        if not rows:
            raise TrainingDataError('Cleaned manifest has no eligible rows')
        return rows
    finally:
        workbook.close()


def _load_reference_samples(reference_root: Path) -> dict[str, tuple[dict, str]]:
    references: dict[str, tuple[dict, str]] = {}
    for path in sorted(reference_root.glob('*.json')):
        content = path.read_bytes()
        payload = _validated_sample(content)
        target = payload['target']
        if target in references:
            raise TrainingDataError(f'Duplicate expert reference for target {target}')
        references[target] = (payload, hashlib.sha256(content).hexdigest())
    return references


def _validated_sample(content: bytes) -> dict:
    # Delayed import prevents a circular dependency when FastAPI exposes model validation.
    from main import Sample
    return Sample.model_validate_json(content).model_dump()


def _path_length(strokes: list[dict]) -> float:
    return sum(stroke['pathLength'] for stroke in strokes)


def _mean(values: list[float], fallback: float) -> float:
    return sum(values) / len(values) if values else fallback


def _feature_measurements(reference: dict, learner: dict) -> dict[str, float]:
    form = compare_form_images(reference['image'], learner['image'], size=128)['measurements']
    reference_strokes, learner_strokes = extract_stroke_features(reference), extract_stroke_features(learner)
    comparison = compare_stroke_sequences(
        [stroke['points'] for stroke in reference_strokes['strokes']],
        [stroke['points'] for stroke in learner_strokes['strokes']],
    )
    reference_paths, learner_paths = reference_strokes['strokes'], learner_strokes['strokes']
    matching_strokes = list(zip(reference_paths, learner_paths))
    stroke_distances = [
        item['normalizedDistance'] for item in comparison['comparisons']
        if 'normalizedDistance' in item
    ]
    endpoint_gaps = [
        (math.dist((left['start']['x'], left['start']['y']), (right['start']['x'], right['start']['y']))
         + math.dist((left['end']['x'], left['end']['y']), (right['end']['x'], right['end']['y']))) / 2
        for left, right in matching_strokes
    ]
    direction_matches = [left['direction'] == right['direction'] for left, right in matching_strokes]
    reference_length, learner_length = _path_length(reference_paths), _path_length(learner_paths)
    if reference_length <= 0:
        raise TrainingDataError('Expert reference has no drawable path length')

    reference_bounds, learner_bounds = form['referenceInkBounds'], form['learnerInkBounds']
    values = {
        'mask_iou': float(form['maskIou']),
        'centroid_distance': float(form['centroidDistance']),
        'ink_width_difference': abs((reference_bounds['right'] - reference_bounds['left']) - (learner_bounds['right'] - learner_bounds['left'])) / 128,
        'ink_height_difference': abs((reference_bounds['bottom'] - reference_bounds['top']) - (learner_bounds['bottom'] - learner_bounds['top'])) / 128,
        'stroke_count_difference': float(abs(reference_strokes['strokeCount'] - learner_strokes['strokeCount'])),
        'stroke_path_distance': _mean(stroke_distances, 1.0),
        'path_length_ratio': learner_length / reference_length,
        'start_end_displacement': _mean(endpoint_gaps, 1.0),
        'direction_match_ratio': _mean([float(match) for match in direction_matches], 0.0),
    }
    if not all(math.isfinite(value) for value in values.values()):
        raise TrainingDataError('Feature extraction produced a non-finite value')
    return {name: round(value, 12) for name, value in values.items()}


def build_feature_rows(manifest_rows: list[dict], project_root: Path, expert_reference_root: Path) -> list[dict]:
    """Create one finite feature row per eligible manifest record."""
    references = _load_reference_samples(expert_reference_root)
    rows = []
    for manifest in manifest_rows:
        sample_path = _inside(project_root, project_root / manifest['relative_path'])
        if not sample_path.is_file():
            raise TrainingDataError(f"Missing sample file: {manifest['relative_path']}")
        content = sample_path.read_bytes()
        sample = _validated_sample(content)
        if sample['target'] != manifest['target']:
            raise TrainingDataError(f"Sample target mismatch for {manifest['sample_id']}")
        if sample['target'] not in references:
            raise TrainingDataError(f"Missing expert reference for target {sample['target']}")
        reference, reference_checksum = references[sample['target']]
        measurements = _feature_measurements(reference, sample)
        rows.append({
            'feature_schema_version': FEATURE_SCHEMA_VERSION,
            'sample_id': manifest['sample_id'],
            'writer_code': manifest['writer_code'],
            'target': manifest['target'],
            'condition': manifest['condition'],
            'binary_label': manifest['binary_label'],
            'sample_checksum': hashlib.sha256(content).hexdigest(),
            'reference_checksum': reference_checksum,
            'feature_names': list(measurements),
            **measurements,
        })
    return rows


def extract_features_for_sample(sample: dict, expert_reference_root: Path) -> dict:
    """Extract one deployable pilot feature row against the matching expert reference."""
    target = sample.get('target')
    references = _load_reference_samples(expert_reference_root)
    if target not in references:
        raise TrainingDataError(f'Missing expert reference for target {target}')
    reference, reference_checksum = references[target]
    measurements = _feature_measurements(reference, sample)
    return {
        'feature_schema_version': FEATURE_SCHEMA_VERSION, 'target': target,
        'reference_checksum': reference_checksum, 'feature_names': list(measurements), **measurements,
    }


def writer_disjoint_folds(rows: list[dict]) -> list[dict]:
    """Return leave-one-writer-out splits without fitting a model."""
    writers = sorted({row['writer_code'] for row in rows})
    if len(writers) < 2:
        raise TrainingDataError('At least two writers are required for writer-disjoint splits')
    return [
        {
            'held_out_writer': writer,
            'train_rows': [row for row in rows if row['writer_code'] != writer],
            'test_rows': [row for row in rows if row['writer_code'] == writer],
        }
        for writer in writers
    ]


def main() -> None:
    parser = argparse.ArgumentParser(description='Extract pilot Baybayin model features from the cleaned Excel manifest.')
    parser.add_argument('--manifest', type=Path, required=True)
    parser.add_argument('--project-root', type=Path, required=True)
    parser.add_argument('--expert-reference-root', type=Path, required=True)
    parser.add_argument('--features-output', type=Path, required=True)
    parser.add_argument('--folds-output', type=Path, required=True)
    args = parser.parse_args()

    manifest_rows = load_cleaned_manifest(args.manifest)
    feature_rows = build_feature_rows(manifest_rows, args.project_root, args.expert_reference_root)
    folds = writer_disjoint_folds(feature_rows)
    args.features_output.parent.mkdir(parents=True, exist_ok=True)
    args.features_output.write_text('\n'.join(json.dumps(row, sort_keys=True) for row in feature_rows) + '\n', encoding='utf-8')
    args.folds_output.parent.mkdir(parents=True, exist_ok=True)
    args.folds_output.write_text(json.dumps([
        {
            'held_out_writer': fold['held_out_writer'],
            'train_sample_ids': [row['sample_id'] for row in fold['train_rows']],
            'test_sample_ids': [row['sample_id'] for row in fold['test_rows']],
        }
        for fold in folds
    ], indent=2) + '\n', encoding='utf-8')


if __name__ == '__main__':
    main()
